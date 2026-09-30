import csv
import io
import os
import pathlib
import tempfile

SUPPORTED = {".md", ".txt", ".csv", ".xlsx", ".xlsm", ".docx", ".pdf"}
MAX_CHARS = 12000


def _xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                out.append(" | ".join(cells).rstrip(" |"))
    return "\n".join(out)


def _docx(path):
    import docx

    d = docx.Document(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(out)


def _pdf(path):
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _text(path):
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _csv(path):
    text = _text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
    except csv.Error:
        return text
    return "\n".join(" | ".join(r) for r in csv.reader(io.StringIO(text), dialect))


READERS = {".xlsx": _xlsx, ".xlsm": _xlsx, ".docx": _docx, ".pdf": _pdf, ".csv": _csv}


def extract_text(path):
    path = pathlib.Path(path)
    try:
        text = READERS.get(path.suffix.lower(), _text)(path)
    except Exception as e:
        text = f"(could not read {path.name}: {e})"
    return text[:MAX_CHARS]


def extract_bytes(name, data):
    suffix = pathlib.Path(name).suffix
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as f:
        f.write(data)
    try:
        return extract_text(f.name)
    finally:
        os.unlink(f.name)


def load_folder(folder):
    folder = pathlib.Path(folder)
    return {
        str(p.relative_to(folder)): extract_text(p)
        for p in sorted(folder.rglob("*"))
        if p.is_file() and p.suffix.lower() in SUPPORTED
    }
